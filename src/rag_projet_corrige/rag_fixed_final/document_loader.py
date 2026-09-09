"""
Chargement et normalisation des documents pour le pipeline RAG TELNET.

Formats supportés :
- Markdown (.md)
- Texte (.txt)
- PDF (.pdf)
- Word (.docx)

Le loader :
- parcourt récursivement le dossier data/
- nettoie légèrement les textes
- préserve la structure
- conserve les métadonnées importantes
- calcule un hash du fichier
- élimine les doublons exacts
"""

from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import List

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

from langchain_community.document_loaders import PyPDFLoader
from langchain_core.documents import Document


class DocumentLoader:
    """
    Charge les documents présents dans un dossier et ses sous-dossiers.
    """

    SUPPORTED_EXTENSIONS = {
        ".md",
        ".txt",
        ".pdf",
        ".docx",
    }

    def __init__(
        self,
        data_dir: str,
        min_chars: int = 30,
    ):
        self.data_dir = Path(data_dir)
        self.min_chars = min_chars

    def load(self) -> List[Document]:
        """
        Charge tous les documents supportés.

        Returns:
            Liste de Documents LangChain.
        """

        self._validate_directory()

        print("=" * 60)
        print("Chargement des documents")
        print("=" * 60)

        # Recherche récursive des fichiers supportés
        files = sorted(
            path
            for path in self.data_dir.rglob("*")
            if path.is_file()
            and path.suffix.lower() in self.SUPPORTED_EXTENSIONS
        )

        print(f"Fichiers détectés : {len(files)}")

        documents: List[Document] = []

        # Ensemble permettant de détecter les doublons
        seen_content_hashes: set[str] = set()

        failed_files = 0
        duplicate_documents = 0

        # Traitement de chaque fichier
        for path in files:

            try:
                loaded_docs = self._load_file(path)

                # Un PDF peut produire plusieurs Documents
                for document in loaded_docs:

                    # Nettoyage du contenu
                    document.page_content = self._clean_text(
                        document.page_content
                    )

                    # Ignore les documents trop courts
                    if len(document.page_content.strip()) < self.min_chars:

                        continue

                    # Hash du contenu nettoyé
                    content_hash = hashlib.sha256(
                        document.page_content.encode("utf-8")
                    ).hexdigest()

                    # Déduplication exacte du contenu
                    if content_hash in seen_content_hashes:

                        duplicate_documents += 1

                        continue

                    seen_content_hashes.add(content_hash)

                    # Hash du fichier original
                    file_hash = self._file_hash(path)

                    # Construction des métadonnées
                    metadata = self._build_metadata(
                        path=path,
                        file_hash=file_hash,
                        content_hash=content_hash,
                        page=document.metadata.get("page"),
                    )

                    # Remplacement des anciennes métadonnées
                    document.metadata = metadata

                    # Ajout du document à la liste finale
                    documents.append(document)

            except Exception as exc:

                failed_files += 1

        print()
        print(f"Documents chargés : {len(documents)}")
        print(f"Doublons ignorés   : {duplicate_documents}")
        print(f"Fichiers en erreur : {failed_files}")

        # Aucun document exploitable
        if not documents:
            raise ValueError(
                f"Aucun document exploitable trouvé dans "
                f"'{self.data_dir}'."
            )

        return documents

    def load_documents(self) -> List[Document]:
        """Alias de compatibilité vers :meth:`load`."""
        return self.load()

    def _validate_directory(self) -> None:
        """
        Vérifie que le dossier existe et est bien un dossier.
        """

        if not self.data_dir.exists():
            raise FileNotFoundError(
                f"Le dossier '{self.data_dir}' n'existe pas."
            )

        if not self.data_dir.is_dir():
            raise NotADirectoryError(
                f"'{self.data_dir}' n'est pas un dossier."
            )

    def _load_file(self, path: Path) -> List[Document]:
        """
        Charge un fichier selon son extension.
        """

        extension = path.suffix.lower()

        # Markdown et texte
        if extension in {".md", ".txt"}:

            text = path.read_text(
                encoding="utf-8",
                errors="replace",
            )

            return [
                Document(
                    page_content=text,
                    metadata={},
                )
            ]

        # PDF
        if extension == ".pdf":

            return PyPDFLoader(str(path)).load()

        # Word
        if extension == ".docx":

            if DocxDocument is None:

                return []

            return [
                self._load_docx(path)
            ]

        raise ValueError(
            f"Format non supporté : {extension}"
        )

    def _load_docx(self, path: Path) -> Document:
        """
        Charge un fichier DOCX avec python-docx.

        Les paragraphes et tableaux sont conservés.
        """

        docx = DocxDocument(path)

        parts: list[str] = []

        # Extraction des paragraphes
        for paragraph in docx.paragraphs:

            text = paragraph.text.strip()

            if text:
                parts.append(text)

        # Extraction des tableaux
        for table in docx.tables:

            for row in table.rows:

                cells = [
                    cell.text.strip().replace("\n", " ")
                    for cell in row.cells
                ]

                row_text = " | ".join(cells)

                if row_text.strip():
                    parts.append(row_text)

        return Document(
            page_content="\n".join(parts),
            metadata={},
        )

    @staticmethod
    def _clean_text(text: str) -> str:
        """
        Nettoyage léger sans détruire la structure.

        On conserve :
        - paragraphes
        - listes
        - titres
        - indentation
        - code
        """

        if not text:
            return ""

        # Suppression du BOM
        text = text.replace("\ufeff", "")

        # Suppression des caractères NULL
        text = text.replace("\x00", "")

        # Normalisation des retours à la ligne
        text = text.replace("\r\n", "\n")
        text = text.replace("\r", "\n")

        # Suppression des espaces en fin de ligne
        lines = [
            line.rstrip()
            for line in text.split("\n")
        ]

        cleaned = "\n".join(lines)

        # Réduction des blocs de lignes vides excessifs
        cleaned = re.sub(
            r"\n{3,}",
            "\n\n",
            cleaned,
        )

        return cleaned.strip()

    def _file_hash(self, path: Path) -> str:
        """
        Calcule un SHA-256 du fichier.

        Ce hash sert notamment à détecter les modifications
        des documents avant une réindexation.
        """

        sha256 = hashlib.sha256()

        with path.open("rb") as file:

            for block in iter(
                lambda: file.read(1024 * 1024),
                b"",
            ):
                sha256.update(block)

        return sha256.hexdigest()

    def _build_metadata(
        self,
        path: Path,
        file_hash: str,
        content_hash: str,
        page: int | None = None,
    ) -> dict:
        """
        Construit les métadonnées standardisées.
        """

        relative_path = path.relative_to(
            self.data_dir
        ).as_posix()

        metadata = {
            "source": str(path),
            "file_path": str(path),
            "file": str(path),
            "filename": path.name,
            "extension": path.suffix.lower(),
            "relative_path": relative_path,
            "file_hash": file_hash,
            "content_hash": content_hash,
        }

        # Ajout du numéro de page pour les PDF
        if page is not None:
            metadata["page"] = int(page)

        # Identifiant stable du document
        if page is not None:
            document_id = f"{file_hash}:page:{page}"
        else:
            document_id = file_hash

        metadata["document_id"] = document_id

        return metadata